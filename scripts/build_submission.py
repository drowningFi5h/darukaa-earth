"""Create the reviewer document from verified local delivery metadata."""

import json
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

root = Path(__file__).resolve().parents[1]
output = root / "output/submission"
output.mkdir(parents=True, exist_ok=True)
metadata = json.loads((root / "tmp/submission-metadata.json").read_text())
doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.7)
section.bottom_margin = Inches(0.65)
section.left_margin = section.right_margin = Inches(0.75)
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(10)
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.line_spacing = 1.08
for name, size in [("Title", 27), ("Heading 1", 15), ("Heading 2", 12)]:
    style = doc.styles[name]
    style.font.name = "Calibri"
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_before = Pt(12)
    style.paragraph_format.space_after = Pt(6)


def link(label, url):
    p = doc.add_paragraph()
    p.add_run(label + " ").bold = True
    rel = p.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    node = OxmlElement("w:hyperlink")
    node.set(qn("r:id"), rel)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "24543A")
    props.append(color)
    run.append(props)
    text = OxmlElement("w:t")
    text.text = url
    run.append(text)
    node.append(run)
    p._p.append(node)


def para(text):
    doc.add_paragraph(text)


def bullet(text):
    doc.add_paragraph(text, style="List Bullet")


doc.add_heading("Darukaa Earth Project Submission", 0)
para("Full stack geospatial platform for carbon and biodiversity projects")
para(
    "Darukaa Earth brings project management, mapped site boundaries, and monthly environmental observations into one workspace. This submission includes a React interface, a Python API, durable PostGIS storage, automated checks, and a reproducible sample dataset."
)
doc.add_heading("Review the application", 1)
link("Public repository", metadata["repo"])
if metadata.get("live_verified"):
    link("Live application", metadata["live_url"])
else:
    para(
        "Public deployment is pending. The application has been verified locally against the Neon PostGIS database; the repository contains complete deployment configuration."
    )
link("Automated checks", metadata["ci_url"])
para(
    "Choose Explore the platform on the home page for the read-only demo. Registration creates a separate workspace where reviewers can create projects and sites."
)
if metadata.get("demo_verified"):
    para("Demo login: demo@darukaa.earth")
    para("Demo password: " + metadata["demo_password"])
para(
    "The demo contains three illustrative Indian project scenarios, six polygons, and twelve monthly observations per site for 2025. These are synthetic examples, not verified environmental benefits or carbon credits."
)
shot = output / "screenshots/workspace-desktop.png"
if shot.exists():
    doc.add_picture(str(shot), width=Inches(5.5))
    doc.paragraphs[-1].alignment = 1
    p = doc.add_paragraph(
        "Landscape explorer with sample project boundaries and calculated site areas."
    )
    p.paragraph_format.space_after = Pt(0)
    p.runs[0].font.size = Pt(9)

doc.add_page_break()
doc.add_heading("Architecture and database", 1)
para(
    "React, TypeScript, Vite, Tailwind CSS, Radix-based UI components, Motion, and TanStack Query power the interface. Mapbox GL JS provides flat maps and polygon drawing; Chart.js displays carbon and biodiversity trends. FastAPI serves both JSON endpoints and the built frontend from one Render Docker service. Neon stores PostgreSQL data with PostGIS."
)
para(
    "Users own projects; projects contain sites; sites have dated measurements. Site polygons use WGS84 geometry with a spatial index. PostGIS computes authoritative hectares using geodesic area. Measurements store monthly carbon removal in tCO2e and observed species counts. A unique site/date constraint prevents duplicate observations."
)
para(
    "Argon2 hashes passwords. JWT sessions use HttpOnly cookies, Secure on HTTPS, with SameSite and Origin checks. Every project/site query enforces ownership. The demo account cannot write. Invalid polygons are rejected and new sites start with no measurements."
)
doc.add_heading("Local setup", 1)
para(
    "Install Node 22.12 or later, Python 3.12, uv, and Docker (or use an existing PostGIS database). Copy .env.example to .env, supply the documented settings, then run the following from the repository root:"
)
for command in [
    "npm ci",
    "uv sync --project backend --frozen",
    "docker compose up -d db",
    "cd backend",
    "uv run alembic upgrade head",
    "uv run python -m app.seed",
    "uv run uvicorn app.main:app --host 127.0.0.1 --port 8000",
]:
    p = doc.add_paragraph(command)
    p.paragraph_format.space_after = Pt(1)
    p.runs[0].font.name = "Consolas"
    p.runs[0].font.size = Pt(9)
para(
    "In another terminal at the repository root, run npm run dev and open http://localhost:5173. For Neon, set DATABASE_URL to its TLS connection string and omit Docker. The README lists every environment variable."
)
doc.add_heading("Quality checks and delivery", 1)
para(
    "Husky and lint-staged enforce formatting and linting before commits. GitHub Actions checks TypeScript/build, ESLint, Prettier, Ruff, and seven integration tests against PostGIS. A successful main-branch run triggers deployment of the tested commit using the authenticated Render deploy API, then verifies that the release becomes live. A deploy hook can also be configured. Render automatic deployments are disabled."
)
para(
    "Verified browser flows include registration, project creation, polygon drawing, save/reload persistence, site editing, demo entry, chart switching, and responsive layouts at 390, 768, and 1440 pixels. Tests also cover ownership isolation, read-only demo access, duplicate registration, invalid polygons, and origin protection."
)
doc.add_heading("Reviewer walkthrough and access", 1)
para(
    "Explore the demo, select a project, and open a site to inspect carbon and biodiversity charts. Use View all landscapes to see every site. Register, create a project, add a site, draw and save a polygon, then reload and edit it. GeoJSON entry is available as a keyboard alternative."
)
para(
    "The repository is public and reviewers can access it directly without invitations. Submit this Word document through the applied-job page using its document-submission option."
)
para(
    "Free Render and Neon services may pause when idle, so initial access can take time. Measurement ingestion, teams, deletion, and password recovery are outside this MVP. The README documents photo credits, data assumptions, security tradeoffs, and deployment details."
)
doc.core_properties.title = "Darukaa Earth Project Submission"
doc.core_properties.subject = "Full stack hackathon application and review instructions"
doc.core_properties.author = "Tausif Ansari"
for style in doc.styles:
    for border in list(style.element.iter(qn("w:pBdr"))):
        border.getparent().remove(border)
for paragraph in doc.paragraphs:
    for border in list(paragraph._p.iter(qn("w:pBdr"))):
        border.getparent().remove(border)
doc.save(output / "Darukaa_Earth_Submission.docx")
print("Created submission Word document")
